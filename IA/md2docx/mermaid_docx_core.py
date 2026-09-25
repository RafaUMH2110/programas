"""
mermaid_docx_core.py
=====================

Núcleo compartido para el pipeline de conversión:

    Markdown (.md) con diagramas Mermaid  --> Word (.docx) en español (es-ES)

Este módulo NO se ejecuta directamente. Es importado tanto por la versión
"VS Code" (``convertir_vscode.py``) como por la versión "CLI"
(``convertir_cli.py``), de modo que la lógica de negocio vive en un único
lugar (principio DRY) y ambos scripts solo se ocupan de la interfaz con el
usuario (configuración, argumentos, mensajes).

Dependencias externas necesarias en el sistema (no son paquetes de Python):
    - Pandoc       -> https://pandoc.org/installing.html
    - Mermaid CLI  -> npm install -g @mermaid-js/mermaid-cli   (comando `mmdc`)

Dependencias de Python (ver requirements.txt):
    - python-docx  (para reforzar el idioma es-ES en el .docx generado)
    - Pillow       (para medir las imágenes PNG y ajustarlas al tamaño de página)

Autor: Generado para un flujo de trabajo de documentación técnica.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
import tempfile
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Sequence, Tuple


# ---------------------------------------------------------------------------
# Constantes y expresiones regulares
# ---------------------------------------------------------------------------

# Detecta bloques ```mermaid ... ``` (multilínea, no codicioso, respeta
# variantes con espacios tras el nombre del lenguaje, p. ej. ```mermaid  ).
PATRON_BLOQUE_MERMAID = re.compile(
    r"```mermaid[ \t]*\r?\n(.*?)```[ \t]*\r?\n?",
    re.DOTALL,
)

TEMAS_VALIDOS = {"default", "dark", "neutral", "forest", "base"}

# Tamaños de página estándar en centímetros (ancho, alto), en orientación vertical.
TAMANOS_PAGINA_CM = {
    "A4": (21.0, 29.7),
    "Carta": (21.59, 27.94),  # "Letter" en Norteamérica
}

# DPI de referencia asumido para convertir píxeles "lógicos" (CSS px, tal y
# como los interpreta el navegador que usa Puppeteer/Mermaid) a centímetros.
# 96 DPI es el estándar de facto para píxeles CSS/web.
DPI_REFERENCIA = 96


# ---------------------------------------------------------------------------
# Excepciones propias
# ---------------------------------------------------------------------------

class HerramientaNoEncontradaError(RuntimeError):
    """Se lanza cuando falta una dependencia externa (Pandoc o Mermaid CLI)."""


class RenderizadoMermaidError(RuntimeError):
    """Se lanza cuando `mmdc` falla al renderizar un diagrama concreto."""


class ConversionPandocError(RuntimeError):
    """Se lanza cuando Pandoc falla al generar el documento .docx."""


# ---------------------------------------------------------------------------
# Estructuras de datos
# ---------------------------------------------------------------------------

@dataclass
class DiagramaMermaid:
    """Representa un diagrama Mermaid localizado dentro del Markdown original."""

    indice: int            # posición secuencial (1, 2, 3...) para nombrar archivos
    codigo: str             # código fuente Mermaid (sin las vallas ```)
    inicio: int              # offset de inicio del bloque completo en el texto
    fin: int                 # offset de fin del bloque completo en el texto
    ruta_png: Optional[Path] = field(default=None)  # se rellena tras renderizar


# ---------------------------------------------------------------------------
# 1. Verificación de dependencias externas
# ---------------------------------------------------------------------------

def verificar_dependencias() -> None:
    """
    Comprueba que Pandoc y Mermaid CLI (mmdc) estén instalados y accesibles
    en el PATH del sistema.

    Lanza:
        HerramientaNoEncontradaError con instrucciones claras de instalación
        si falta alguna herramienta.
    """
    faltantes = []

    if shutil.which("pandoc") is None:
        faltantes.append(
            "  - Pandoc no encontrado.\n"
            "    Instalación:\n"
            "      Windows : winget install --id JohnMacFarlane.Pandoc\n"
            "      macOS   : brew install pandoc\n"
            "      Linux   : sudo apt install pandoc   (o consulta https://pandoc.org/installing.html)"
        )

    if shutil.which("mmdc") is None:
        faltantes.append(
            "  - Mermaid CLI (mmdc) no encontrado.\n"
            "    Instalación (requiere Node.js >= 16):\n"
            "      npm install -g @mermaid-js/mermaid-cli\n"
            "    Verifica con: mmdc --version"
        )

    if faltantes:
        mensaje = (
            "No se pudieron localizar una o más herramientas externas necesarias:\n\n"
            + "\n\n".join(faltantes)
            + "\n\nInstala las herramientas indicadas y vuelve a ejecutar el script."
        )
        raise HerramientaNoEncontradaError(mensaje)


# ---------------------------------------------------------------------------
# 2. Extracción de diagramas Mermaid del Markdown
# ---------------------------------------------------------------------------

def extraer_diagramas_mermaid(contenido_md: str) -> List[DiagramaMermaid]:
    """
    Busca todos los bloques ```mermaid ... ``` dentro del texto Markdown.

    Args:
        contenido_md: contenido íntegro del archivo Markdown de entrada.

    Returns:
        Lista de objetos DiagramaMermaid en el orden en que aparecen en el
        documento (índice 1, 2, 3...).
    """
    diagramas: List[DiagramaMermaid] = []
    for i, coincidencia in enumerate(PATRON_BLOQUE_MERMAID.finditer(contenido_md), start=1):
        codigo = coincidencia.group(1).strip("\n")
        diagramas.append(
            DiagramaMermaid(
                indice=i,
                codigo=codigo,
                inicio=coincidencia.start(),
                fin=coincidencia.end(),
            )
        )
    return diagramas


# ---------------------------------------------------------------------------
# 3. Renderizado de diagramas a PNG mediante Mermaid CLI
# ---------------------------------------------------------------------------

def renderizar_diagrama(
    diagrama: DiagramaMermaid,
    directorio_salida: Path,
    tema: str = "default",
    ancho: int = 1200,
    alto: int = 800,
    escala: float = 2.0,
    fondo: str = "white",
    sin_sandbox: bool = False,
    verbose: bool = False,
) -> Path:
    """
    Renderiza un único diagrama Mermaid como imagen PNG de alta calidad
    usando el comando `mmdc` (Mermaid CLI).

    Args:
        diagrama: objeto DiagramaMermaid con el código fuente a renderizar.
        directorio_salida: carpeta donde se guardarán los archivos temporales
            (.mmd) y la imagen PNG resultante.
        tema: tema visual de Mermaid ("default", "dark", "neutral", "forest", "base").
        ancho: ancho del lienzo en píxeles.
        alto: alto del lienzo en píxeles.
        escala: factor de escala para aumentar la resolución final (nitidez).
        fondo: color de fondo ("white", "transparent", "#RRGGBB", etc.).
        sin_sandbox: si True, añade --no-sandbox a Puppeteer (útil en
            contenedores Docker/CI de Linux donde el sandbox de Chromium falla).
        verbose: si True, muestra la salida completa de `mmdc`.

    Returns:
        Ruta al archivo PNG generado.

    Lanza:
        RenderizadoMermaidError si `mmdc` termina con código de error.
    """
    directorio_salida.mkdir(parents=True, exist_ok=True)

    nombre_base = f"diagrama_{diagrama.indice:03d}_{uuid.uuid4().hex[:8]}"
    ruta_mmd = directorio_salida / f"{nombre_base}.mmd"
    ruta_png = directorio_salida / f"{nombre_base}.png"

    # Escribimos el código Mermaid a un archivo temporal (mmdc requiere un archivo, no stdin fiable)
    ruta_mmd.write_text(diagrama.codigo, encoding="utf-8")

    comando = [
        "mmdc",
        "-i", str(ruta_mmd),
        "-o", str(ruta_png),
        "-t", tema,
        "-w", str(ancho),
        "-H", str(alto),
        "-s", str(escala),
        "-b", fondo,
    ]

    ruta_puppeteer_cfg: Optional[Path] = None
    if sin_sandbox:
        # Generamos una configuración temporal de Puppeteer para desactivar el sandbox.
        ruta_puppeteer_cfg = directorio_salida / f"{nombre_base}_puppeteer.json"
        ruta_puppeteer_cfg.write_text(
            '{"args": ["--no-sandbox", "--disable-setuid-sandbox"]}', encoding="utf-8"
        )
        comando.extend(["-p", str(ruta_puppeteer_cfg)])

    if verbose:
        print(f"    → Ejecutando: {' '.join(comando)}")

    try:
        resultado = subprocess.run(
            comando,
            capture_output=True,
            text=True,
            check=False,
        )
    except FileNotFoundError as exc:
        raise RenderizadoMermaidError(
            "No se pudo ejecutar 'mmdc'. ¿Está Mermaid CLI instalado y en el PATH?"
        ) from exc
    finally:
        # El .mmd temporal ya no es necesario, independientemente del resultado.
        ruta_mmd.unlink(missing_ok=True)
        if ruta_puppeteer_cfg is not None:
            ruta_puppeteer_cfg.unlink(missing_ok=True)

    if resultado.returncode != 0 or not ruta_png.exists():
        raise RenderizadoMermaidError(
            f"Fallo al renderizar el diagrama #{diagrama.indice}.\n"
            f"Salida de mmdc (stderr):\n{resultado.stderr.strip()}"
        )

    if verbose:
        print(f"    ✓ Diagrama #{diagrama.indice} renderizado en: {ruta_png}")

    return ruta_png


def renderizar_todos_los_diagramas(
    diagramas: Sequence[DiagramaMermaid],
    directorio_salida: Path,
    tema: str = "default",
    ancho: int = 1200,
    alto: int = 800,
    escala: float = 2.0,
    fondo: str = "white",
    sin_sandbox: bool = False,
    verbose: bool = False,
) -> None:
    """
    Renderiza en orden todos los diagramas detectados, actualizando cada
    objeto DiagramaMermaid con la ruta de su imagen PNG (campo `ruta_png`).
    """
    total = len(diagramas)
    for diagrama in diagramas:
        print(f"  [{diagrama.indice}/{total}] Renderizando diagrama Mermaid...")
        diagrama.ruta_png = renderizar_diagrama(
            diagrama,
            directorio_salida,
            tema=tema,
            ancho=ancho,
            alto=alto,
            escala=escala,
            fondo=fondo,
            sin_sandbox=sin_sandbox,
            verbose=verbose,
        )


# ---------------------------------------------------------------------------
# 4. Ajuste del tamaño de las imágenes al área útil de una página
# ---------------------------------------------------------------------------

def _normalizar_tamano_pagina(tamano_pagina: str) -> str:
    """
    Normaliza el nombre de un tamaño de página para aceptar variantes
    habituales ("a4", "A4", "carta", "Letter"...) y devuelve la clave
    canónica usada en TAMANOS_PAGINA_CM.

    Lanza:
        ValueError si el tamaño de página no es reconocido.
    """
    clave = tamano_pagina.strip()
    equivalencias = {"letter": "Carta", "carta": "Carta", "a4": "A4"}
    clave_normalizada = equivalencias.get(clave.lower(), clave)

    if clave_normalizada not in TAMANOS_PAGINA_CM:
        raise ValueError(
            f"Tamaño de página no reconocido: '{tamano_pagina}'. "
            f"Valores admitidos: {', '.join(TAMANOS_PAGINA_CM)} (o 'Letter')."
        )
    return clave_normalizada


def calcular_area_util_pagina(
    tamano_pagina: str = "A4",
    margen_cm: float = 2.5,
) -> Tuple[float, float]:
    """
    Calcula el área útil (ancho, alto) en centímetros de una página, es
    decir, el tamaño de página menos los márgenes izquierdo/derecho y
    superior/inferior. Se usa como límite máximo para las imágenes
    insertadas, de modo que ningún diagrama supere el tamaño de una página.

    Args:
        tamano_pagina: "A4" o "Carta" (Letter). No distingue mayúsculas/minúsculas.
        margen_cm: margen aplicado en cada lado, en centímetros.

    Returns:
        Tupla (ancho_util_cm, alto_util_cm).

    Lanza:
        ValueError si `tamano_pagina` no es reconocido o el margen es excesivo.
    """
    clave_normalizada = _normalizar_tamano_pagina(tamano_pagina)
    ancho_pagina_cm, alto_pagina_cm = TAMANOS_PAGINA_CM[clave_normalizada]
    ancho_util_cm = ancho_pagina_cm - (2 * margen_cm)
    alto_util_cm = alto_pagina_cm - (2 * margen_cm)

    if ancho_util_cm <= 0 or alto_util_cm <= 0:
        raise ValueError(
            f"El margen indicado ({margen_cm} cm) es demasiado grande para "
            f"el tamaño de página '{clave_normalizada}'."
        )

    return ancho_util_cm, alto_util_cm


def medir_imagen_en_cm(
    ruta_png: Path,
    escala: float = 2.0,
    dpi_referencia: int = DPI_REFERENCIA,
) -> Tuple[float, float]:
    """
    Mide las dimensiones reales de un PNG ya renderizado y las convierte a
    centímetros "lógicos" (independientes de la sobremuestreada aplicada).

    Mermaid CLI (`mmdc`) no siempre respeta al pie de la letra el ancho/alto
    solicitados (`-w`/`-H`): ajusta el lienzo al contenido real del diagrama.
    Por eso medimos el PNG resultante en lugar de asumir los valores pedidos.

    El factor `escala` (parámetro `-s` de mmdc) sobremuestrea la imagen para
    obtener más nitidez (más píxeles por la misma área visual), así que se
    divide por ese factor antes de convertir a centímetros, o el diagrama
    parecería más grande de lo que realmente se verá.

    Args:
        ruta_png: ruta al archivo PNG ya renderizado.
        escala: factor de sobremuestreo usado al renderizar (`-s` de mmdc).
        dpi_referencia: DPI asumido para la conversión de píxeles a cm.

    Returns:
        Tupla (ancho_cm, alto_cm) con el tamaño "lógico" de la imagen.

    Lanza:
        HerramientaNoEncontradaError si Pillow no está instalado.
    """
    try:
        from PIL import Image
    except ImportError as exc:
        raise HerramientaNoEncontradaError(
            "Falta la librería 'Pillow', necesaria para medir las imágenes "
            "y ajustarlas al tamaño de página.\n"
            "Instálala con: pip install Pillow"
        ) from exc

    with Image.open(ruta_png) as imagen:
        ancho_px, alto_px = imagen.size

    escala = escala if escala > 0 else 1.0
    ancho_logico_px = ancho_px / escala
    alto_logico_px = alto_px / escala

    ancho_cm = (ancho_logico_px / dpi_referencia) * 2.54
    alto_cm = (alto_logico_px / dpi_referencia) * 2.54
    return ancho_cm, alto_cm


def ajustar_dimensiones_a_maximo(
    ancho_cm: float,
    alto_cm: float,
    ancho_max_cm: float,
    alto_max_cm: float,
) -> Tuple[float, float]:
    """
    Calcula el tamaño final (ancho, alto) en centímetros de una imagen de
    modo que quepa dentro de los límites máximos indicados, preservando la
    relación de aspecto original. Nunca amplía una imagen que ya sea más
    pequeña que el límite (factor de escalado nunca superior a 1.0).

    Args:
        ancho_cm, alto_cm: dimensiones originales de la imagen.
        ancho_max_cm, alto_max_cm: límites máximos permitidos (p. ej. el
            área útil de una página).

    Returns:
        Tupla (ancho_final_cm, alto_final_cm), siempre <= los máximos dados.
    """
    if ancho_cm <= 0 or alto_cm <= 0:
        return ancho_cm, alto_cm

    factor = min(ancho_max_cm / ancho_cm, alto_max_cm / alto_cm, 1.0)
    return ancho_cm * factor, alto_cm * factor


def preparar_referencia_con_geometria(
    directorio_temporal: Path,
    tamano_pagina: str = "A4",
    margen_cm: float = 2.5,
    verbose: bool = False,
) -> Optional[Path]:
    """
    Genera una plantilla `--reference-doc` temporal con el tamaño de página
    y los márgenes indicados, partiendo de la plantilla de referencia por
    defecto de Pandoc (para conservar sus estilos de título, cuerpo, etc.).

    Es necesario porque, SIN una plantilla de referencia explícita, Pandoc
    aplica un límite interno de ancho de imagen (~14.8 cm) que no coincide
    necesariamente con el área útil real calculada a partir del tamaño de
    página y margen solicitados por el usuario. Al fijar la geometría de
    página en la propia plantilla de referencia, Pandoc respeta el ancho
    exacto que le indiquemos en cada imagen (hasta el límite del área útil
    real de esa página).

    Args:
        directorio_temporal: carpeta temporal donde crear el archivo.
        tamano_pagina: "A4" o "Carta".
        margen_cm: margen aplicado en cada lado, en centímetros.
        verbose: si True, informa del proceso.

    Returns:
        Ruta a la plantilla de referencia generada, o None si no se pudo
        generar (en cuyo caso se usará el comportamiento por defecto de
        Pandoc, con su límite interno habitual).
    """
    clave_normalizada = _normalizar_tamano_pagina(tamano_pagina)
    ancho_pagina_cm, alto_pagina_cm = TAMANOS_PAGINA_CM[clave_normalizada]

    ruta_referencia = directorio_temporal / "referencia_geometria.docx"

    resultado = subprocess.run(
        ["pandoc", "-o", str(ruta_referencia), "--print-default-data-file", "reference.docx"],
        capture_output=True,
        text=True,
        check=False,
    )
    if resultado.returncode != 0 or not ruta_referencia.exists():
        if verbose:
            print(
                "  ⚠ No se pudo generar una plantilla de referencia con la geometría de "
                "página deseada; Pandoc usará su límite de ancho de imagen por defecto."
            )
        return None

    try:
        from docx import Document
        from docx.shared import Cm

        documento = Document(str(ruta_referencia))
        for seccion in documento.sections:
            seccion.page_width = Cm(ancho_pagina_cm)
            seccion.page_height = Cm(alto_pagina_cm)
            seccion.left_margin = Cm(margen_cm)
            seccion.right_margin = Cm(margen_cm)
            seccion.top_margin = Cm(margen_cm)
            seccion.bottom_margin = Cm(margen_cm)
        documento.save(str(ruta_referencia))
    except Exception as exc:  # noqa: BLE001 — no bloquea el resto del proceso
        if verbose:
            print(f"  ⚠ No se pudo ajustar la geometría de la plantilla de referencia: {exc}")
        return None

    if verbose:
        print(
            f"  ✓ Plantilla de referencia generada con página {clave_normalizada} "
            f"({ancho_pagina_cm}x{alto_pagina_cm} cm, margen {margen_cm} cm): {ruta_referencia}"
        )

    return ruta_referencia


# ---------------------------------------------------------------------------
# 5. Sustitución de los bloques Mermaid por sintaxis de imagen Markdown
# ---------------------------------------------------------------------------

def sustituir_diagramas_por_imagenes(
    contenido_md: str,
    diagramas: Sequence[DiagramaMermaid],
    escala: float = 2.0,
    tamano_pagina: str = "A4",
    margen_cm: float = 2.5,
    verbose: bool = False,
) -> str:
    """
    Reemplaza cada bloque ```mermaid ... ``` del texto original por una
    referencia de imagen Markdown (![...](ruta.png){width=...cm}) que
    apunta al PNG ya renderizado, conservando el resto del documento intacto.

    Antes de insertar cada imagen, se mide su tamaño real y se calcula un
    ancho ajustado (atributo `width` de Pandoc, en centímetros) para que
    NINGÚN diagrama supere el área útil de una página del documento final.
    Al indicar solo el `width`, Pandoc conserva automáticamente la relación
    de aspecto al calcular la altura.

    Se procesa en orden inverso (de atrás hacia adelante) para que los
    offsets (inicio/fin) de los bloques anteriores no se vean invalidados
    por los cambios de longitud del texto ya sustituido.

    Args:
        contenido_md: Markdown original (con los bloques mermaid intactos).
        diagramas: lista de DiagramaMermaid ya renderizados (con ruta_png).
        escala: factor de sobremuestreo usado al renderizar los PNG.
        tamano_pagina: "A4" o "Carta", usado para calcular el límite máximo.
        margen_cm: margen de página (en cm) usado para calcular el área útil.
        verbose: si True, informa del tamaño original y del ajustado.

    Returns:
        Nuevo contenido Markdown con las imágenes insertadas en las
        posiciones exactas donde estaban los bloques de código, ya
        ajustadas para no exceder el tamaño de una página.
    """
    ancho_max_cm, alto_max_cm = calcular_area_util_pagina(tamano_pagina, margen_cm)

    nuevo_contenido = contenido_md
    for diagrama in sorted(diagramas, key=lambda d: d.inicio, reverse=True):
        if diagrama.ruta_png is None:
            raise ValueError(
                f"El diagrama #{diagrama.indice} no ha sido renderizado todavía."
            )

        ancho_cm, alto_cm = medir_imagen_en_cm(diagrama.ruta_png, escala=escala)
        ancho_final_cm, alto_final_cm = ajustar_dimensiones_a_maximo(
            ancho_cm, alto_cm, ancho_max_cm, alto_max_cm
        )

        if verbose:
            if round(ancho_final_cm, 2) < round(ancho_cm, 2):
                print(
                    f"    ↳ Diagrama #{diagrama.indice}: "
                    f"{ancho_cm:.1f}x{alto_cm:.1f} cm → reducido a "
                    f"{ancho_final_cm:.1f}x{alto_final_cm:.1f} cm "
                    f"(límite de página: {ancho_max_cm:.1f}x{alto_max_cm:.1f} cm)"
                )
            else:
                print(
                    f"    ↳ Diagrama #{diagrama.indice}: {ancho_cm:.1f}x{alto_cm:.1f} cm "
                    "(ya cabe dentro de una página, sin cambios)"
                )

        # Usamos ruta POSIX (barras /) porque Pandoc la interpreta igual en
        # todos los sistemas operativos y evita problemas de escape en Windows.
        ruta_imagen = diagrama.ruta_png.as_posix()
        # Solo se especifica 'width': Pandoc calcula la altura automáticamente
        # para conservar la relación de aspecto original de la imagen.
        etiqueta = (
            f"![Diagrama {diagrama.indice}]({ruta_imagen})"
            f"{{width={ancho_final_cm:.2f}cm}}"
        )
        nuevo_contenido = (
            nuevo_contenido[: diagrama.inicio]
            + etiqueta
            + "\n"
            + nuevo_contenido[diagrama.fin :]
        )
    return nuevo_contenido


# ---------------------------------------------------------------------------
# 6. Metadatos de idioma para Pandoc (es-ES)
# ---------------------------------------------------------------------------

def preparar_markdown_con_metadatos(
    contenido_md: str,
    idioma: str = "es-ES",
    titulo: Optional[str] = None,
) -> str:
    """
    Antepone un bloque de metadatos YAML al Markdown para que Pandoc
    configure el idioma del documento .docx resultante (afecta a
    corrector ortográfico, guionización y formato de fecha/hora nativo
    de Word cuando el documento se abre en español).

    Args:
        contenido_md: Markdown (ya con las imágenes sustituidas).
        idioma: código de idioma BCP-47, por defecto "es-ES".
        titulo: título opcional del documento (metadato `title`).

    Returns:
        Markdown con el front-matter YAML añadido al principio.
    """
    lineas_yaml = ["---", f"lang: {idioma}"]
    if titulo:
        # Escapamos comillas dobles por seguridad dentro del YAML.
        titulo_escapado = titulo.replace('"', '\\"')
        lineas_yaml.append(f'title: "{titulo_escapado}"')
    lineas_yaml.append("---\n")
    return "\n".join(lineas_yaml) + "\n" + contenido_md


# ---------------------------------------------------------------------------
# 7. Conversión final a .docx mediante Pandoc
# ---------------------------------------------------------------------------

def convertir_a_docx(
    ruta_md_temporal: Path,
    ruta_docx_salida: Path,
    referencia_docx: Optional[Path] = None,
    directorio_recursos: Optional[Path] = None,
    verbose: bool = False,
) -> None:
    """
    Invoca a Pandoc para convertir el Markdown (ya procesado, con imágenes
    y metadatos de idioma) al documento Word final.

    Args:
        ruta_md_temporal: Markdown intermedio, listo para convertir.
        ruta_docx_salida: ruta del archivo .docx final.
        referencia_docx: plantilla opcional (--reference-doc) para aplicar
            estilos corporativos/editoriales predefinidos.
        directorio_recursos: carpeta adicional donde buscar imágenes
            referenciadas con rutas relativas (--resource-path).
        verbose: si True, muestra el comando y la salida de Pandoc.

    Lanza:
        ConversionPandocError si Pandoc devuelve un código de error.
    """
    ruta_docx_salida.parent.mkdir(parents=True, exist_ok=True)

    comando = [
        "pandoc",
        str(ruta_md_temporal),
        "-o", str(ruta_docx_salida),
        "--standalone",
        "--from", "markdown+yaml_metadata_block",
    ]

    if directorio_recursos is not None:
        comando.extend(["--resource-path", str(directorio_recursos)])

    if referencia_docx is not None:
        if not referencia_docx.exists():
            raise FileNotFoundError(
                f"La plantilla de referencia no existe: {referencia_docx}"
            )
        comando.extend(["--reference-doc", str(referencia_docx)])

    if verbose:
        print(f"  → Ejecutando: {' '.join(comando)}")

    resultado = subprocess.run(comando, capture_output=True, text=True, check=False)

    if resultado.returncode != 0:
        raise ConversionPandocError(
            f"Pandoc finalizó con errores:\n{resultado.stderr.strip()}"
        )

    if verbose and resultado.stderr:
        print(f"  (avisos de Pandoc): {resultado.stderr.strip()}")


# ---------------------------------------------------------------------------
# 8. Refuerzo del idioma es-ES directamente en el .docx (post-procesado)
# ---------------------------------------------------------------------------

def forzar_idioma_docx(ruta_docx: Path, idioma: str = "es-ES") -> None:
    """
    Refuerza el idioma español en el documento .docx generado, actuando en
    dos niveles complementarios al metadato `lang` ya aplicado por Pandoc:

      1. Propiedades del documento (core properties) -> idioma del propio archivo.
      2. Marca de idioma (w:lang) en cada `run` de texto -> asegura que el
         corrector ortográfico/gramatical de Word trate todo el contenido
         como español, incluso si algún párrafo no heredó el idioma del
         estilo "Normal" definido por Pandoc.

    Requiere la librería `python-docx`.

    Args:
        ruta_docx: ruta al archivo .docx ya generado por Pandoc.
        idioma: código de idioma BCP-47 (por defecto "es-ES").
    """
    from docx import Document
    from docx.oxml.ns import qn

    documento = Document(str(ruta_docx))

    # 1) Propiedad de idioma a nivel de documento.
    try:
        documento.core_properties.language = idioma
    except Exception:
        # Algunas versiones de python-docx no exponen 'language'; no es crítico.
        pass

    # 2) Marca w:lang en cada run de cada párrafo (incluye tablas).
    def _marcar_idioma_en_parrafo(parrafo) -> None:
        for run in parrafo.runs:
            rPr = run._element.get_or_add_rPr()
            lang_el = rPr.find(qn("w:lang"))
            if lang_el is None:
                lang_el = rPr.makeelement(qn("w:lang"), {})
                rPr.append(lang_el)
            lang_el.set(qn("w:val"), idioma)
            lang_el.set(qn("w:eastAsia"), idioma)
            lang_el.set(qn("w:bidi"), idioma)

    for parrafo in documento.paragraphs:
        _marcar_idioma_en_parrafo(parrafo)

    for tabla in documento.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for parrafo in celda.paragraphs:
                    _marcar_idioma_en_parrafo(parrafo)

    documento.save(str(ruta_docx))


# ---------------------------------------------------------------------------
# 9. Limpieza de archivos temporales
# ---------------------------------------------------------------------------

def limpiar_temporales(rutas: Sequence[Path], verbose: bool = False) -> None:
    """
    Elimina archivos y/o carpetas temporales generados durante el proceso.
    Ignora silenciosamente rutas que ya no existan.
    """
    for ruta in rutas:
        try:
            if ruta.is_dir():
                shutil.rmtree(ruta, ignore_errors=True)
            elif ruta.exists():
                ruta.unlink()
            if verbose:
                print(f"  🧹 Eliminado temporal: {ruta}")
        except OSError as exc:
            print(f"  ⚠ No se pudo eliminar '{ruta}': {exc}", file=sys.stderr)


# ---------------------------------------------------------------------------
# 10. Orquestador de alto nivel (usado por ambos scripts)
# ---------------------------------------------------------------------------

def procesar_markdown_a_docx(
    ruta_md: Path,
    ruta_docx: Path,
    tema: str = "default",
    ancho: int = 1200,
    alto: int = 800,
    escala: float = 2.0,
    fondo: str = "white",
    idioma: str = "es-ES",
    tamano_pagina: str = "A4",
    margen_cm: float = 2.5,
    referencia_docx: Optional[Path] = None,
    mantener_temporales: bool = False,
    sin_sandbox: bool = False,
    verbose: bool = False,
) -> Path:
    """
    Ejecuta el pipeline completo:

        1. Verifica dependencias externas (Pandoc, mmdc).
        2. Lee el Markdown de entrada.
        3. Extrae los diagramas Mermaid.
        4. Renderiza cada diagrama como PNG de alta resolución.
        5. Mide cada PNG y calcula un ancho ajustado para que ningún
           diagrama supere el área útil de una página (A4/Carta).
        6. Sustituye los bloques de código por referencias de imagen
           (con el ancho ya ajustado).
        7. Añade metadatos de idioma (es-ES) al Markdown intermedio.
        8. Convierte con Pandoc al .docx final.
        9. Refuerza el idioma es-ES directamente en el .docx.
        10. Limpia los archivos temporales (salvo que se indique lo contrario).

    Args:
        ruta_md: archivo Markdown de entrada.
        ruta_docx: archivo .docx de salida deseado.
        tema: tema visual de los diagramas Mermaid.
        ancho, alto, escala, fondo: parámetros de renderizado de imágenes.
        idioma: código de idioma para el documento final.
        tamano_pagina: "A4" o "Carta"; define el límite máximo de tamaño
            que puede ocupar cada diagrama insertado.
        margen_cm: margen de página (cm) usado para calcular el área útil
            disponible para las imágenes.
        referencia_docx: plantilla Word opcional para aplicar estilos.
        mantener_temporales: si True, no se borran los PNG ni el .md intermedio.
        sin_sandbox: desactiva el sandbox de Chromium/Puppeteer (Docker/CI Linux).
        verbose: activa mensajes detallados de depuración.

    Returns:
        Ruta final del documento .docx generado.
    """
    if not ruta_md.exists():
        raise FileNotFoundError(f"El archivo Markdown de entrada no existe: {ruta_md}")

    print("→ Verificando dependencias externas (Pandoc, Mermaid CLI)...")
    verificar_dependencias()
    print("  ✓ Todas las dependencias están disponibles.\n")

    print(f"→ Leyendo Markdown de entrada: {ruta_md}")
    contenido_md = ruta_md.read_text(encoding="utf-8")

    diagramas = extraer_diagramas_mermaid(contenido_md)
    print(f"  ✓ Se encontraron {len(diagramas)} diagrama(s) Mermaid.\n")

    # Directorio temporal para PNGs y Markdown intermedio.
    directorio_temporal = Path(tempfile.mkdtemp(prefix="mermaid_docx_"))

    try:
        if diagramas:
            print("→ Renderizando diagramas Mermaid a PNG...")
            renderizar_todos_los_diagramas(
                diagramas,
                directorio_temporal,
                tema=tema,
                ancho=ancho,
                alto=alto,
                escala=escala,
                fondo=fondo,
                sin_sandbox=sin_sandbox,
                verbose=verbose,
            )
            print()

        print("→ Sustituyendo bloques Mermaid por imágenes en el Markdown...")
        print(f"  (ajustando tamaño máximo a una página {tamano_pagina}, margen {margen_cm} cm)")
        contenido_procesado = sustituir_diagramas_por_imagenes(
            contenido_md,
            diagramas,
            escala=escala,
            tamano_pagina=tamano_pagina,
            margen_cm=margen_cm,
            verbose=verbose,
        )

        print(f"→ Aplicando metadatos de idioma ({idioma})...")
        contenido_final = preparar_markdown_con_metadatos(
            contenido_procesado, idioma=idioma, titulo=ruta_md.stem
        )

        ruta_md_temporal = directorio_temporal / "documento_intermedio.md"
        ruta_md_temporal.write_text(contenido_final, encoding="utf-8")

        # Si el usuario no ha aportado su propia plantilla, generamos una
        # temporal con la geometría de página exacta (tamaño + márgenes).
        # Esto es necesario porque, sin ninguna plantilla de referencia,
        # Pandoc aplica un límite interno de ancho de imagen (~14.8 cm) que
        # no tiene por qué coincidir con el área útil real solicitada.
        referencia_efectiva = referencia_docx
        if referencia_efectiva is None:
            print(f"→ Preparando geometría de página ({tamano_pagina}, margen {margen_cm} cm)...")
            referencia_efectiva = preparar_referencia_con_geometria(
                directorio_temporal, tamano_pagina=tamano_pagina, margen_cm=margen_cm, verbose=verbose
            )

        print(f"→ Convirtiendo a Word con Pandoc: {ruta_docx}")
        convertir_a_docx(
            ruta_md_temporal,
            ruta_docx,
            referencia_docx=referencia_efectiva,
            directorio_recursos=directorio_temporal,
            verbose=verbose,
        )
        print("  ✓ Documento .docx generado correctamente.\n")

        print(f"→ Reforzando idioma {idioma} en el documento final...")
        forzar_idioma_docx(ruta_docx, idioma=idioma)
        print("  ✓ Idioma aplicado.\n")

    finally:
        if mantener_temporales:
            print(f"ℹ Archivos temporales conservados en: {directorio_temporal}")
        else:
            limpiar_temporales([directorio_temporal], verbose=verbose)

    return ruta_docx
